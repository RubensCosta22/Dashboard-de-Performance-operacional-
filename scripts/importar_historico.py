"""
Importador de Histórico — Demo
"""

import argparse, re, sys, datetime
from pathlib import Path

# Adiciona raiz do projeto ao path
sys.path.insert(0, str(Path(__file__).parent.parent))

from openpyxl import load_workbook

# ── CONFIG ─────────────────────────────────────────────────────
PERIODOS = [
    ('01X07', 'p0'), ('07X13', 'p1'), ('13X19', 'p2'), ('19X01', 'p3')
]

IGNORAR = {'', 'SELECIONE', 'Total Parado:', 'FARELO', 'SOJA', 'MILHO'}

# ── CORREÇÃO PRINCIPAL AQUI 🔥 ──────────────────────────────────
def extrair_num_dia(nome_arquivo):
    """Extrai o dia do nome do arquivo em vários formatos"""

    # (15)
    m = re.search(r'\((\d+)\)', nome_arquivo)
    if m:
        return int(m.group(1))

    # 01.xlsm
    m = re.search(r'^(\d{1,2})', nome_arquivo)
    if m:
        return int(m.group(1))

    # qualquer número no nome
    m = re.search(r'(\d{1,2})', nome_arquivo)
    if m:
        return int(m.group(1))

    return None


def t2m(val):
    if val is None: return 0
    if isinstance(val, datetime.timedelta):
        return int(val.total_seconds() / 60)
    if isinstance(val, datetime.time):
        return val.hour * 60 + val.minute
    s = str(val).strip()
    m = re.match(r'(\d+):(\d+)', s)
    if m:
        return int(m.group(1)) * 60 + int(m.group(2))
    return 0


def extrair_arquivo(path, ano, mes, dia):
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception as e:
        print(f"  ⚠ Erro ao abrir {path.name}: {e}")
        return None

    resultado = {'data': f"{ano}-{mes:02d}-{dia:02d}", 'periodos': {}}

    for aba_nome, pid in PERIODOS:
        if aba_nome not in wb.sheetnames:
            continue

        ws = wb[aba_nome]
        rows = list(ws.iter_rows(max_row=130, values_only=True))

        periodo = {'paradas': [], 'descarga': []}

        for row in rows:
            if not row or len(row) < 10:
                continue

            causa = str(row[3]).strip() if row[3] else ''
            if causa in IGNORAR:
                continue

            dur = t2m(row[8])
            if dur == 0:
                continue

            periodo['paradas'].append({
                'causa': causa,
                'duracao': dur
            })

        resultado['periodos'][pid] = periodo

    return resultado


def gravar_no_banco(dados_lista, excel_path):
    from openpyxl import load_workbook

    if not excel_path.exists():
        print("⚠ Banco não encontrado. Rode o sistema primeiro.")
        return 0

    wb = load_workbook(excel_path)

    ws_t = wb['Turnos']
    ws_d = wb['Descarga']
    ws_p = wb['Paradas']

    def ultimo_id(ws):
        return max((r[0] for r in ws.iter_rows(min_row=2, values_only=True) if r[0]), default=0)

    now = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    gravados = 0

    for dados in dados_lista:
        data = dados['data']

        for pid, periodo in dados['periodos'].items():

            if not periodo['paradas']:
                continue

            turno_id = ultimo_id(ws_t) + 1

            ws_t.append([
                turno_id,
                data,
                pid,
                '', '', '', '', '',
                now
            ])

            # Paradas
            par_id = ultimo_id(ws_p)
            for i, p in enumerate(periodo['paradas'], 1):
                ws_p.append([
                    par_id + i,
                    turno_id,
                    'm1',
                    p['causa'],
                    'term',
                    '',
                    '',
                    p['duracao'],
                    ''
                ])

            gravados += 1

    wb.save(excel_path)
    return gravados
def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--pasta', required=True)
    parser.add_argument('--mes', type=int, required=True)
    parser.add_argument('--ano', type=int, required=True)
    args = parser.parse_args()

    pasta = Path(args.pasta)

    if not pasta.exists():
        print("❌ Pasta não encontrada")
        sys.exit(1)

    excel_path = Path(__file__).parent.parent / 'data' / 'relatorios.xlsx'

    arquivos = sorted(
        pasta.glob('*.xlsm'),
        key=lambda f: extrair_num_dia(f.name) or 0
    )

    print("\n========================================")
    print(f"Importando {len(arquivos)} arquivos")
    print("========================================\n")

    todos = []

    for arq in arquivos:
        dia = extrair_num_dia(arq.name)

        if not dia:
            print(f"⚠ Nome inválido: {arq.name}")
            continue

        print(f"Lendo {arq.name} (dia {dia})...")

        dados = extrair_arquivo(arq, args.ano, args.mes, dia)

        if dados:
            todos.append(dados)

    print("\nGravando no banco...")
    gravados = gravar_no_banco(todos, excel_path)

    print(f"✅ {gravados} registros inseridos\n")


if __name__ == "__main__":
    main()